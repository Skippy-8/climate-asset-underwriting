"""Workbook helpers shared by both models: styles, the Inputs and Sources tabs, time-series rows
and defined names.

Conventions (stated on each Cover tab):
* blue on pale yellow = input; black = formula on the same sheet; green = formula that reads
  another sheet or a named input;
* every input lives once on the Inputs tab under a defined name, so formulas read
  ``=Units*Rent_Growth`` rather than ``=150*0.03`` and contain no hidden constants;
* time series run left to right, one column per year, starting in column D.
"""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from cau.inputs import Input, fingerprint, load_sources

FIRST_COL = 4                       # column D = year 1
INK = "1F2933"
BLUE = Font(name="Arial", size=10, color="0000FF")
BLACK = Font(name="Arial", size=10, color="000000")
GREEN = Font(name="Arial", size=10, color="006100")
BOLD = Font(name="Arial", size=10, bold=True, color=INK)
TITLE = Font(name="Arial", size=14, bold=True, color=INK)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
MUTED = Font(name="Arial", size=9, italic=True, color="616E7C")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
HEAD_FILL = PatternFill("solid", fgColor="334E68")
BAND_FILL = PatternFill("solid", fgColor="F0F4F8")
TAG_FILL = {"Sourced": "DFF3E4", "Estimated": "FFF1D6", "Illustrative": "E9E6F5", "Definition": "EEF1F4"}
THIN = Side(style="thin", color="BCCCDC")
TOP = Border(top=THIN)

FMT = {
    "usd": '#,##0;(#,##0);"–"', "usd2": '#,##0.00;(#,##0.00);"–"', "num": '#,##0;(#,##0);"–"',
    "num1": '#,##0.0', "num2": '#,##0.00', "pct": '0.00%', "pct1": '0.0%', "x": '0.00"x"',
    "int": '0', "text": '@', "sci": '0.0E+0', "num4": '#,##0.0000',
}


def col(i: int) -> str:
    return get_column_letter(i)


def year_cols(n: int) -> list[str]:
    return [col(FIRST_COL + k) for k in range(n)]


def git_commit(repo: Path) -> str:
    try:
        sha = subprocess.run(["git", "-C", str(repo), "rev-parse", "--short", "HEAD"], capture_output=True,
                             text=True, timeout=5).stdout.strip()
        dirty = subprocess.run(["git", "-C", str(repo), "status", "--porcelain", "cau", "data"],
                               capture_output=True, text=True, timeout=5).stdout.strip()
        return (sha or "uncommitted") + ("+local changes" if dirty and sha else "")
    except Exception:
        return "unknown"


class Book:
    """A workbook plus the bookkeeping that keeps formulas readable: named inputs, named outputs,
    and a row index per sheet so formulas can reference rows by name."""

    def __init__(self, n_years: int):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.n = n_years
        self.cols = year_cols(n_years)
        self.rows: dict[str, dict[str, int]] = {}
        self.names: dict[str, str] = {}

    # ── names ───────────────────────────────────────────────────────────────
    def name(self, name: str, sheet: str, ref: str) -> None:
        if name in self.names:
            raise ValueError(f"defined name {name} used twice")
        self.names[name] = f"'{sheet}'!{ref}" if " " in sheet else f"{sheet}!{ref}"
        self.wb.defined_names[name] = DefinedName(name, attr_text=self.names[name])

    def r(self, sheet: str, row: str) -> int:
        return self.rows[sheet][row]

    def ref(self, sheet: str, row: str, c: str, here: str | None = None) -> str:
        """A reference to row ``row`` of ``sheet`` in column ``c``; sheet-qualified unless on ``here``."""
        cell = f"{c}{self.r(sheet, row)}"
        return cell if sheet == here else f"{sheet}!{cell}"

    def span(self, sheet: str, row: str, here: str | None = None, first: int = 0, last: int | None = None) -> str:
        last = self.n - 1 if last is None else last
        a, b = self.cols[first], self.cols[last]
        rr = self.r(sheet, row)
        rng = f"{a}{rr}:{b}{rr}"
        return rng if sheet == here else f"{sheet}!{rng}"

    # ── sheets ──────────────────────────────────────────────────────────────
    def sheet(self, title: str, heading: str, note: str = "", year_header: bool = True, year0: bool = False):
        ws = self.wb.create_sheet(title)
        self.rows[title] = {}
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 15
        for c in self.cols:
            ws.column_dimensions[c].width = 13
        ws["A1"] = heading
        ws["A1"].font = TITLE
        if note:
            ws["A2"] = note
            ws["A2"].font = MUTED
        if year_header:
            ws["A4"] = "Year"
            ws["A4"].font = HEAD
            ws["B4"] = "Unit"
            ws["B4"].font = HEAD
            for c in ["A", "B", "C", *self.cols]:
                ws[f"{c}4"].fill = HEAD_FILL
                ws[f"{c}4"].font = HEAD
                ws[f"{c}4"].alignment = Alignment(horizontal="center" if c not in "AB" else "left")
            if year0:
                ws["C4"] = 0
            first = self.cols[0]
            ws[f"{first}4"] = 1
            for prev, c in zip(self.cols, self.cols[1:], strict=False):
                ws[f"{c}4"] = f"={prev}4+1"
            self.rows[title]["Year"] = 4
        ws.freeze_panes = "D5" if year_header else None
        return ws

    def section(self, ws, row: int, text: str) -> None:
        ws[f"A{row}"] = text
        ws[f"A{row}"].font = BOLD
        for c in ["A", "B", "C", *self.cols]:
            ws[f"{c}{row}"].fill = BAND_FILL

    def series(self, ws, row: int, key: str, label: str, unit: str, fn, fmt: str = "usd", total: bool = False,
               name_range: bool = True) -> None:
        """A time-series row: ``fn(c, prev, k)`` returns the formula for column ``c`` (year k+1)."""
        self.rows[ws.title][key] = row
        ws[f"A{row}"] = label
        ws[f"B{row}"] = unit
        ws[f"B{row}"].font = MUTED
        prev = None
        for k, c in enumerate(self.cols):
            f = fn(c, prev, k)
            cell = ws[f"{c}{row}"]
            cell.value = f
            cell.number_format = FMT[fmt]
            cell.font = GREEN if isinstance(f, str) and _reads_elsewhere(f) else BLACK
            prev = c
        if total:
            for c in ["A", "B", *self.cols]:
                ws[f"{c}{row}"].border = TOP
            ws[f"A{row}"].font = BOLD
        if name_range:
            self.name(f"R_{key}", ws.title, f"${self.cols[0]}${row}:${self.cols[-1]}${row}")

    def scalar(self, ws, row: int, key: str, label: str, unit: str, formula, fmt: str = "usd",
               bold: bool = False, define: bool = True) -> None:
        self.rows[ws.title][key] = row
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = BOLD if bold else BLACK
        ws[f"B{row}"] = unit
        ws[f"B{row}"].font = MUTED
        cell = ws[f"C{row}"]
        cell.value = formula
        cell.number_format = FMT[fmt]
        cell.font = GREEN if isinstance(formula, str) and _reads_elsewhere(formula) else BLACK
        if bold:
            cell.font = Font(name="Arial", size=10, bold=True, color=cell.font.color)
        if define:
            self.name(key, ws.title, f"$C${row}")

    # ── standard tabs ───────────────────────────────────────────────────────
    def inputs_sheet(self, inputs: list[Input], scenarios: dict[str, list[float]] | None = None) -> None:
        """``scenarios`` maps an input key to its [Base, Downside, Severe] values. Those inputs get the
        three values in columns I:K and a live value =CHOOSE(Case_Select, I, J, K) in column C, so one
        switch moves the whole workbook between cases."""
        scenarios = scenarios or {}
        ws = self.wb.create_sheet("Inputs")
        self.rows["Inputs"] = {}
        ws.sheet_view.showGridLines = False
        for c, w in zip("ABCDEFGHIJK", (22, 52, 14, 16, 13, 24, 60, 3, 12, 12, 12), strict=True):
            ws.column_dimensions[c].width = w
        ws["A1"] = "Inputs"
        ws["A1"].font = TITLE
        ws["A2"] = ("Every assumption in the model, once. Tag: Sourced = read from the cited public source; "
                    "Estimated = analyst assumption anchored to a source; Illustrative = deal definition; "
                    "Definition = unit conversion.")
        ws["A2"].font = MUTED
        hdr = ["Section", "Input", "Value", "Unit", "Tag", "Source id", "Note"]
        if scenarios:
            hdr = [*hdr, "", "Base", "Downside", "Severe"]
        for c, h in zip("ABCDEFGHIJK", hdr, strict=False):
            ws[f"{c}4"] = h
            ws[f"{c}4"].font = HEAD
            ws[f"{c}4"].fill = HEAD_FILL
        if scenarios:
            ws["I3"] = "Stress cases: the live value in column C follows Case_Select (1 = Base, 2 = Downside, 3 = Severe)"
            ws["I3"].font = MUTED
        row = 5
        for i in inputs:
            ws[f"A{row}"], ws[f"B{row}"], ws[f"C{row}"] = i.section, i.label, i.value
            ws[f"D{row}"], ws[f"E{row}"], ws[f"F{row}"], ws[f"G{row}"] = i.unit, i.tag, i.source, i.note
            ws[f"C{row}"].font = BLUE
            ws[f"C{row}"].fill = INPUT_FILL
            ws[f"C{row}"].number_format = _input_fmt(i)
            ws[f"E{row}"].fill = PatternFill("solid", fgColor=TAG_FILL[i.tag])
            if i.key in scenarios:
                cases = scenarios[i.key]
                if cases[0] != i.value:
                    raise ValueError(f"{i.key}: scenario Base value {cases[0]} differs from the input {i.value}")
                for c, val in zip("IJK", cases, strict=True):
                    ws[f"{c}{row}"] = val
                    ws[f"{c}{row}"].font = BLUE
                    ws[f"{c}{row}"].fill = INPUT_FILL
                    ws[f"{c}{row}"].number_format = _input_fmt(i)
                ws[f"C{row}"] = f"=CHOOSE(Case_Select,I{row},J{row},K{row})"
                ws[f"C{row}"].font = GREEN
                ws[f"C{row}"].fill = PatternFill(fill_type=None)
            for c in "ABDFG":
                ws[f"{c}{row}"].font = MUTED if c in "DFG" else BLACK
            self.rows["Inputs"][i.key] = row
            self.name(i.key, "Inputs", f"$C${row}")
            row += 1
        ws.freeze_panes = "C5"

    def sources_sheet(self, inputs: list[Input]) -> None:
        ws = self.wb.create_sheet("Sources")
        ws.sheet_view.showGridLines = False
        for c, w in zip("ABCD", (24, 110, 70, 12), strict=True):
            ws.column_dimensions[c].width = w
        ws["A1"] = "Sources"
        ws["A1"].font = TITLE
        for c, h in zip("ABCD", ["Source id", "Citation", "URL", "Retrieved"], strict=True):
            ws[f"{c}3"] = h
            ws[f"{c}3"].font = HEAD
            ws[f"{c}3"].fill = HEAD_FILL
        used = {i.source for i in inputs}
        for row, (sid, s) in enumerate(((k, v) for k, v in load_sources().items() if k in used), start=4):
            ws[f"A{row}"], ws[f"B{row}"], ws[f"C{row}"], ws[f"D{row}"] = sid, s["citation"], s["url"], str(s["retrieved"])
            ws[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")

    def cover(self, title: str, subtitle: str, inputs: list[Input], tabs: list[tuple[str, str]], repo: Path,
              checks_name: str) -> None:
        ws = self.wb.create_sheet("Cover", 0)
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 100
        ws["A1"] = title
        ws["A1"].font = Font(name="Arial", size=16, bold=True, color=INK)
        ws["A2"] = subtitle
        ws["A2"].font = MUTED
        sources = load_sources()
        retrieved = max(str(sources[i.source]["retrieved"]) for i in inputs)
        meta = [("Status", "Illustrative underwriting case built on public benchmarks. Not a real property, "
                           "transaction or offer; no proprietary data."),
                ("Model checks", f'=IF({checks_name},"All checks pass","CHECK FAILURES - see Checks tab")'),
                ("Code version", git_commit(repo)),
                ("Inputs fingerprint", fingerprint(inputs)),
                ("Data retrieved", retrieved),
                ("Colour convention", "Blue on yellow = input (Inputs tab only). Green = formula reading another "
                                      "tab or a named input. Black = formula on the same tab."),
                ("Periods", "Annual. Year 1 starts at closing.")]
        for row, (k, v) in enumerate(meta, start=4):
            ws[f"A{row}"], ws[f"B{row}"] = k, v
            ws[f"A{row}"].font = BOLD
            ws[f"B{row}"].alignment = Alignment(wrap_text=True)
        row = 4 + len(meta) + 1
        ws[f"A{row}"] = "Tabs"
        ws[f"A{row}"].font = BOLD
        for name, desc in tabs:
            row += 1
            ws[f"A{row}"], ws[f"B{row}"] = name, desc

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.calculation.fullCalcOnLoad = True      # Excel recalculates every formula on open
        self.wb.save(path)
        return path


_A1 = re.compile(r"^\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?$")


def _reads_elsewhere(formula: str) -> bool:
    """True when a formula reads another tab or a named input (shown green), false when it only
    reads cells on its own tab (shown black)."""
    if "!" in formula:
        return True
    for tok in Tokenizer(formula).items:
        if tok.type == Token.OPERAND and tok.subtype == Token.RANGE and not _A1.match(tok.value):
            return True
    return False


def _input_fmt(i: Input) -> str:
    if "%" in i.unit:
        return "0.00%"
    if i.unit in ("x",):
        return '0.00"x"'
    if i.value and abs(i.value) < 0.001:
        return "0.000000"
    if i.unit.startswith("$") and abs(i.value) < 100:
        return "#,##0.00"
    if abs(i.value) >= 1000:
        return "#,##0"
    return "0.####"
