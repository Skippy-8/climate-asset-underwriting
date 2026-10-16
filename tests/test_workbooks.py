"""Both Excel workbooks: live formulas, no hidden constants, and the same numbers as their engines."""
import numpy as np
import pytest
from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Token, Tokenizer

from cau.inputs import fingerprint

# Named workbook cells that are deliberately not engine outputs (layout helpers and indices).
WORKBOOK_ONLY = {"Units", "GPR_Y1", "NRSF", "Rate_Fixed", "Rate_Float", "Binding_Test", "NOI_Fwd_SQ", "NOI_Fwd_RF",
                 "Lev_IRR_Delta", "BESS_Capex_Base", "R_Rent_Index", "R_OI_Index", "R_Exp_Index", "R_Tax_Index",
                 "R_Util_Index", "R_Shared_Opex", "R_Monthly_Payment", "R_DS_Check", "R_Cov_In_Hold", "R_In_Hold",
                 "R_Exit_Flag", "R_PV_On", "R_BESS_On", "R_Deg_Index", "R_PPA_On", "R_Toll_On", "R_PPA_Price_Row",
                 "R_Merchant_Price_Row", "R_Esc_Index", "R_Tenor_On", "R_DSRA_Release", "R_Opex_PV", "R_Opex_BESS",
                 "R_Opex_AM", "R_Price_Base", "R_Price_Op", "R_Rev_PV_P50", "R_Rev_PV_P90", "R_Rev_PV_Op"}


def _formulas(path):
    for ws in load_workbook(path):
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    yield ws.title, c.coordinate, c.value


def _same(got, want) -> bool:
    if isinstance(want, (bool, np.bool_)):
        return bool(got) == bool(want)
    return got is not None and abs(float(got) - want) <= 1e-6 * max(1.0, abs(want))


def _rows_match(got, want) -> bool:
    if want.dtype.kind in "US":
        return list(got) == list(want)
    mask = ~np.isnan(want)                       # engine NaN = blank cell in the workbook (outside the tenor)
    return np.allclose(np.asarray(got, dtype=object)[mask].astype(float), want[mask], rtol=1e-6, atol=1e-5)


def test_calculation_tabs_hold_formulas_not_typed_numbers(model):
    wb = load_workbook(model["xlsx"])
    for name in model["calc_sheets"]:
        for row in wb[name].iter_rows(min_row=5):
            for c in row:
                assert not isinstance(c.value, (int, float)) or isinstance(c.value, bool), \
                    f"{name}!{c.coordinate} holds a typed number {c.value!r}; inputs belong on the Inputs tab"


def test_formulas_contain_no_magic_numbers(model):
    offenders = [f"{s}!{a}: {f}" for s, a, f in _formulas(model["xlsx"])
                 for tok in Tokenizer(f).items
                 if tok.type == Token.OPERAND and tok.subtype == Token.NUMBER and tok.value not in {"0", "1"}]
    assert not offenders, "constants inside formulas:\n" + "\n".join(offenders[:20])


def test_formulas_are_well_formed(model):
    for sheet, coord, f in _formulas(model["xlsx"]):
        Tokenizer(f)
        assert f.count("(") == f.count(")"), f"{sheet}!{coord}"


def test_every_named_output_matches_the_engine(model):
    res, got = model["result"], model["eval"]
    bad = [f"{k}: engine {v} workbook {got[k]}" for k, v in res.scalars.items() if not _same(got[k], v)]
    bad += [f"row {k}" for k, v in res.rows.items() if f"R_{k}" in got and not _rows_match(got[f"R_{k}"], v)]
    bad += [f"cash flow {k}" for k, v in res.returns.items() if not _rows_match(got[f"R_{k}"], v)]
    missing = [k for k in res.scalars if k not in got]
    known = set(res.scalars) | {f"R_{k}" for k in [*res.rows, *res.returns]} | {i.key for i in model["inputs"]}
    unchecked = [n for n in got if n not in known and not n.startswith(("Check_", "All_Checks", "MACRS_", "Snap_"))
                 and n not in WORKBOOK_ONLY]
    assert not bad and not missing and not unchecked, "\n".join(
        bad + [f"missing name {m}" for m in missing] + [f"workbook output not in the engine: {u}" for u in unchecked])


def test_all_engine_rows_are_named_in_the_workbook(model):
    got = model["eval"]
    unnamed = [k for k in model["result"].rows if f"R_{k}" not in got]
    assert not unnamed, unnamed


def test_workbook_checks_all_pass(model):
    got = model["eval"]
    failing = [k for k, v in got.items() if k.startswith("Check_") and v is not True]
    assert not failing and got["All_Checks"] is True, failing


def test_inputs_tab_is_complete_and_tagged(model):
    wb = load_workbook(model["xlsx"])
    ws = wb["Inputs"]
    named = {n for n, d in wb.defined_names.items() if d.attr_text.startswith("Inputs!") and n != "MACRS_Table"}
    assert named == {i.key for i in model["inputs"]}
    tags = {ws[f"E{r}"].value for r in range(5, 5 + len(model["inputs"]))}
    assert tags <= {"Sourced", "Estimated", "Illustrative", "Definition"}


def test_committed_workbook_is_current_and_carries_values(model):
    """The committed file must be rebuilt from today's inputs and, after the Excel recalculation in
    scripts/build_models.py, carry stored values that match the engine."""
    path = model["committed"]
    cover = load_workbook(path)["Cover"]
    stamp = {cover[f"A{r}"].value: cover[f"B{r}"].value for r in range(4, 12)}
    assert stamp["Inputs fingerprint"] == fingerprint(model["inputs"]), "rebuild models/ (scripts/build_models.py)"
    vals, names = load_workbook(path, data_only=True), load_workbook(path).defined_names
    sheet, ref = names["All_Checks"].attr_text.split("!")
    if vals[sheet][ref.replace("$", "")].value is None:
        pytest.skip("committed workbook has no stored values (built without Excel)")
    errors = [f"{ws.title}!{c.coordinate}" for ws in vals for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#")]
    assert not errors, errors[:10]
    for k, want in model["result"].scalars.items():
        sheet, ref = names[k].attr_text.split("!")
        assert _same(vals[sheet][ref.replace("$", "")].value, want), k
